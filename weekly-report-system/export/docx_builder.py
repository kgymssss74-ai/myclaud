"""
docx_builder.py — Generate weekly report DOCX.

TG order: Audio Solution → Audio SW → Wireless Audio
Delta spans stored in delta_spans table → blue runs (RGB 0,0,255).
Images embedded from item_attachments paths.
"""
import io
import sqlite3
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

TG_ORDER = ['Audio Solution', 'Audio SW', 'Wireless Audio']
BLUE = RGBColor(0, 0, 255)

FIELD_LABELS = {
    'progress_text': '진행',
    'schedule_text': '일정',
    'risk_text':     'Risk',
}


def _conn(db_path: str) -> sqlite3.Connection:
    c = sqlite3.connect(db_path)
    c.row_factory = sqlite3.Row
    return c


def _write_with_delta(para, text: str, spans: list[tuple[int, int]]) -> None:
    """Add runs to paragraph; spans in curr_text are blue."""
    if not text:
        return
    if not spans:
        para.add_run(text)
        return
    last = 0
    for s, e in spans:
        if s > last:
            para.add_run(text[last:s])
        run = para.add_run(text[s:e])
        run.font.color.rgb = BLUE
        last = e
    if last < len(text):
        para.add_run(text[last:])


def _milestone_table(doc: Document, conn: sqlite3.Connection, projects) -> None:
    project_ids = [p['project_id'] for p in projects]
    if not project_ids:
        return

    doc.add_heading('마일스톤 현황', level=2)
    tbl = doc.add_table(rows=1, cols=6)
    tbl.style = 'Table Grid'
    hdr = tbl.rows[0].cells
    for i, h in enumerate(['과제', '유형', 'MS', '계획일', '실적일', '상태']):
        hdr[i].text = h

    placeholders = ','.join('?' * len(project_ids))
    rows = conn.execute(
        f'''SELECT p.name proj, p.track, m.type, m.planned_date, m.actual_date, m.status
            FROM milestones m
            JOIN projects p ON m.project_id = p.project_id
            WHERE m.project_id IN ({placeholders})
            ORDER BY p.name, m.ms_id''',
        project_ids,
    ).fetchall()

    for r in rows:
        cells = tbl.add_row().cells
        cells[0].text = r['proj']
        cells[1].text = r['track']
        cells[2].text = r['type']
        cells[3].text = r['planned_date'] or ''
        cells[4].text = r['actual_date'] or ''
        cells[5].text = r['status']
        if r['status'] == 'risk':
            for cell in cells:
                for para in cell.paragraphs:
                    for run in para.runs:
                        run.font.color.rgb = RGBColor(0xFF, 0, 0)


def build_doc(week: str, db_path: str) -> bytes:
    conn = _conn(db_path)
    doc = Document()
    doc.add_heading(f'주간보고 {week}', level=0)

    # Collect all reported projects for milestone table
    all_projects = conn.execute(
        '''SELECT DISTINCT p.project_id, p.name, p.track
           FROM weekly_reports r
           JOIN projects p ON r.project_id = p.project_id
           WHERE r.week = ?''',
        (week,),
    ).fetchall()
    _milestone_table(doc, conn, all_projects)
    doc.add_paragraph()

    for tg in TG_ORDER:
        members = conn.execute(
            'SELECT member_id, display_name, name FROM members WHERE tg=? ORDER BY tg_order, member_id',
            (tg,),
        ).fetchall()

        tg_heading_added = False
        for member in members:
            reports = conn.execute(
                '''SELECT r.report_id, p.project_id, p.name project_name, p.track
                   FROM weekly_reports r
                   JOIN projects p ON r.project_id = p.project_id
                   WHERE r.week=? AND r.member_id=?
                   ORDER BY p.name''',
                (week, member['member_id']),
            ).fetchall()
            if not reports:
                continue

            if not tg_heading_added:
                doc.add_heading(tg, level=1)
                tg_heading_added = True

            display = member['display_name'] or member['name']
            doc.add_heading(display, level=2)

            for report in reports:
                doc.add_heading(report['project_name'], level=3)
                items = conn.execute(
                    'SELECT * FROM report_items WHERE report_id=? ORDER BY item_id',
                    (report['report_id'],),
                ).fetchall()

                for item in items:
                    is_new = item['prev_item_id'] is None
                    key_label = item['item_key']
                    if is_new:
                        key_label += ' [신규]'
                    doc.add_paragraph(key_label, style='Heading 4')

                    spans_by_field = {}
                    for span in conn.execute(
                        'SELECT field, start_pos, end_pos FROM delta_spans WHERE item_id=? ORDER BY start_pos',
                        (item['item_id'],),
                    ).fetchall():
                        spans_by_field.setdefault(span['field'], []).append(
                            (span['start_pos'], span['end_pos'])
                        )

                    for field, label in FIELD_LABELS.items():
                        text = item[field] or ''
                        para = doc.add_paragraph()
                        para.add_run(f'{label}: ').bold = True
                        _write_with_delta(para, text, spans_by_field.get(field, []))

                    if item['delay_reason']:
                        doc.add_paragraph(f'지연사유: {item["delay_reason"]}')

                    # Attach images
                    for att in conn.execute(
                        'SELECT path FROM item_attachments WHERE item_id=? AND kind=?',
                        (item['item_id'], 'image'),
                    ).fetchall():
                        img_path = Path(db_path).parent / att['path']
                        if img_path.exists():
                            try:
                                doc.add_picture(str(img_path), width=Inches(5))
                            except Exception:
                                doc.add_paragraph(f'[이미지: {att["path"]}]')

    conn.close()
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
