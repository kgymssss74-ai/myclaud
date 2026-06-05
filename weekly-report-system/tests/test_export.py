"""Gate ⑦⑧ — DOC / XLS 생성 검증"""
import io
import zipfile

import openpyxl
import pytest
from docx import Document

from export.docx_builder import build_doc
from export.xlsx_builder import build_eval_xls, build_feedback_xls, build_mbo_xls


# ── DOC 생성 ──────────────────────────────────────────────────────────────────

def test_docx_is_valid_zip(app_with_report):
    with app_with_report.app_context():
        data = build_doc('2026-W23', app_with_report.config['DB_PATH'])
    assert isinstance(data, bytes)
    assert zipfile.is_zipfile(io.BytesIO(data)), 'DOCX must be a valid ZIP/OOXML file'


def test_docx_is_openable(app_with_report):
    with app_with_report.app_context():
        data = build_doc('2026-W23', app_with_report.config['DB_PATH'])
    doc = Document(io.BytesIO(data))
    assert doc is not None


def test_docx_contains_tg_heading(app_with_report):
    with app_with_report.app_context():
        data = build_doc('2026-W23', app_with_report.config['DB_PATH'])
    doc = Document(io.BytesIO(data))
    headings = [p.text for p in doc.paragraphs if p.style.name.startswith('Heading')]
    # Should have at least one heading referencing Audio Solution TG
    assert any('Audio Solution' in h for h in headings), f'Headings: {headings}'


def test_docx_has_member_section(app_with_report):
    with app_with_report.app_context():
        data = build_doc('2026-W23', app_with_report.config['DB_PATH'])
    doc = Document(io.BytesIO(data))
    all_text = '\n'.join(p.text for p in doc.paragraphs)
    assert '김강열' in all_text


def test_docx_has_blue_delta_run(app_with_report):
    """Delta spans are stored for PVR item → run at that span should be blue."""
    from docx.shared import RGBColor
    with app_with_report.app_context():
        data = build_doc('2026-W23', app_with_report.config['DB_PATH'])
    doc = Document(io.BytesIO(data))
    blue = RGBColor(0, 0, 255)
    blue_runs = [
        run for para in doc.paragraphs
        for run in para.runs
        if run.font.color.rgb == blue
    ]
    assert len(blue_runs) > 0, 'Expected at least one blue run for delta span'


def test_docx_milestone_table_present(app_with_report):
    with app_with_report.app_context():
        data = build_doc('2026-W23', app_with_report.config['DB_PATH'])
    doc = Document(io.BytesIO(data))
    assert len(doc.tables) >= 1, 'Expected milestone table in DOC'


def test_docx_route_returns_docx(app_with_report):
    client = app_with_report.test_client()
    resp = client.get('/export/doc?week=2026-W23')
    assert resp.status_code == 200
    ct = resp.content_type
    assert 'wordprocessingml' in ct or 'openxmlformats' in ct


# ── MBO XLS ───────────────────────────────────────────────────────────────────

def test_mbo_xls_is_valid(app_with_report):
    with app_with_report.app_context():
        data = build_mbo_xls(2026, app_with_report.config['DB_PATH'])
    assert isinstance(data, bytes)
    wb = openpyxl.load_workbook(io.BytesIO(data))
    assert wb is not None


def test_mbo_xls_has_achievement_formula(app_with_report):
    with app_with_report.app_context():
        data = build_mbo_xls(2026, app_with_report.config['DB_PATH'])
    wb = openpyxl.load_workbook(io.BytesIO(data), data_only=False)
    ws = wb.active
    # Header row: col G should be '달성률(%)'
    headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
    assert '달성률(%)' in headers
    g_col = headers.index('달성률(%)') + 1
    # Data row 2 should have a formula in col G
    formula = ws.cell(2, g_col).value
    assert formula is not None and str(formula).startswith('='), f'Expected formula, got: {formula}'


def test_mbo_xls_route_returns_xlsx(app_with_report):
    client = app_with_report.test_client()
    resp = client.get('/export/xls/mbo?year=2026')
    assert resp.status_code == 200
    assert 'spreadsheet' in resp.content_type or 'openxmlformats' in resp.content_type


# ── Feedback XLS ──────────────────────────────────────────────────────────────

def test_feedback_xls_is_valid(app_with_report):
    with app_with_report.app_context():
        data = build_feedback_xls('', '', app_with_report.config['DB_PATH'])
    wb = openpyxl.load_workbook(io.BytesIO(data))
    ws = wb.active
    headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
    assert '상태' in headers
    assert '진행' in headers


def test_feedback_xls_has_data_row(app_with_report):
    with app_with_report.app_context():
        data = build_feedback_xls('2026-W23', '2026-W23', app_with_report.config['DB_PATH'])
    wb = openpyxl.load_workbook(io.BytesIO(data))
    ws = wb.active
    assert ws.max_row >= 2, f'Expected data rows, max_row={ws.max_row}'


# ── Eval XLS ──────────────────────────────────────────────────────────────────

def test_eval_xls_has_completion_formula(app_with_report):
    with app_with_report.app_context():
        data = build_eval_xls(2026, app_with_report.config['DB_PATH'])
    wb = openpyxl.load_workbook(io.BytesIO(data), data_only=False)
    ws = wb.active
    headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
    assert '완료율(%)' in headers
    rate_col = headers.index('완료율(%)') + 1
    formula = ws.cell(2, rate_col).value
    assert formula is not None and str(formula).startswith('='), f'Expected formula, got: {formula}'


# ── Backup ────────────────────────────────────────────────────────────────────

def test_backup_creates_zip(tmp_path):
    import sqlite3
    # Create a minimal DB
    db_path = str(tmp_path / 'weekly.db')
    conn = sqlite3.connect(db_path)
    conn.execute('CREATE TABLE t (id INTEGER PRIMARY KEY)')
    conn.commit()
    conn.close()

    from backup.backup import run_backup
    out = str(tmp_path / 'backup')
    zip_path = run_backup(db_path, str(tmp_path / 'uploads'), out)

    import os
    assert os.path.exists(zip_path)
    assert zipfile.is_zipfile(zip_path)
